import csv
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from .store import Store, uid


FIELDS = {"name", "supplier_id", "product_id", "purchase_order_id", "delivery_date", "required_date", "quantity", "stock", "owner", "status"}
KINDS = {"Supplier", "Product", "PurchaseOrder", "SalesOrder", "Team", "BusinessRecord"}


def cell(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return "" if value is None else str(value).strip()


def parse_file(filename: str, content: bytes) -> dict:
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("Upload limit is 10 MB")
    suffix = Path(filename).suffix.lower()
    tables, sections, warnings = [], [], []
    if suffix in (".xlsx", ".docx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 3000 or sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Compressed document exceeds processing limits")
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        if not rows:
            raise ValueError("The CSV is empty")
        tables.append({"name": Path(filename).stem, "headers": rows[0], "rows": rows[1:]})
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                if sheet.max_row > 5001 or sheet.max_column > 100:
                    raise ValueError("Workbooks are limited to 5,000 data rows and 100 columns per sheet")
                rows = [[cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
                if rows:
                    tables.append({"name": sheet.title, "headers": rows[0], "rows": rows[1:]})
        finally:
            workbook.close()
        warnings.append("Formula cells use saved cached values. Workweave does not recalculate workbooks or execute macros.")
    elif suffix in (".txt", ".md"):
        text = content.decode("utf-8-sig")
        sections = [{"location": f"Section {i + 1}", "text": s.strip()} for i, s in enumerate(re.split(r"\n\s*\n", text)) if s.strip()]
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 150:
            raise ValueError("Encrypted PDFs or PDFs over 150 pages are unsupported")
        sections = [{"location": f"Page {i + 1}", "text": page.extract_text() or ""} for i, page in enumerate(reader.pages)]
        if not any(s["text"].strip() for s in sections):
            raise ValueError("No text found. Scanned PDFs need OCR, which is not available in this release.")
    elif suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        sections = [{"location": f"Paragraph {i + 1}", "text": p.text} for i, p in enumerate(document.paragraphs) if p.text.strip()]
        for i, table in enumerate(document.tables):
            sections.append({"location": f"Document table {i + 1}", "text": "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)})
    else:
        raise ValueError("Supported files: XLSX, CSV, TXT, MD, text PDF and DOCX")
    if sum(len(t["rows"]) for t in tables) > 5000 or sum(len(s["text"]) for s in sections) > 500000:
        raise ValueError("Import exceeds 5,000 rows or 500,000 text characters")
    for table in tables:
        table["headers"] = [cell(h) for h in table["headers"]]
        if not all(table["headers"]) or len(set(table["headers"])) != len(table["headers"]):
            raise ValueError("Every column needs a unique non-empty heading")
        table["rows"] = [[cell(v) for v in r] for r in table["rows"] if any(cell(v) for v in r)]
        if len(table["headers"]) > 100 or any(len(row) != len(table["headers"]) for row in table["rows"]):
            raise ValueError("Tables need at most 100 columns and consistent row widths")
    return {"tables": tables, "sections": sections, "warnings": warnings}


def suggestions(table: dict) -> dict:
    headers = table["headers"]
    norm = {h.lower().replace("_", "").replace(" ", ""): h for h in headers}
    name = table["name"].lower()
    kind = next((k for k in ("PurchaseOrder", "SalesOrder", "Supplier", "Product", "Team") if k.lower() in name.replace("_", "").replace(" ", "")), "BusinessRecord")
    if kind == "BusinessRecord":
        kind = "SalesOrder" if "required_date" in headers else "PurchaseOrder" if "delivery_date" in headers else kind
    key = next((norm[x] for x in ("id", "sku", "suppliercode", "salesorderid", "purchaseorderid") if x in norm), headers[0])
    fields = {f: norm[f.replace("_", "")] for f in FIELDS if f.replace("_", "") in norm}
    return {"table": table["name"], "entity_type": kind, "key_column": key, "fields": fields}


def chunks(text: str):
    # Approximate 500 tokens by words; retain section boundaries and 80-token overlap.
    words = text.split()
    for start in range(0, len(words), 315):
        yield " ".join(words[start:start + 375])
        if start + 375 >= len(words):
            break


async def preview(store: Store, filename: str, content: bytes, source_key: str | None = None):
    try:
        parsed = parse_file(filename, content)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not parse this file. Check that it is a valid, unencrypted supported document; text and CSV files must use UTF-8.") from exc
    job = {"id": uid("import"), "filename": Path(filename).name, "version": hashlib.sha256(content).hexdigest()[:16],
           "source_key": source_key or Path(filename).name, "status": "preview", **parsed}
    await store.put("import", job)
    return public_preview(job)


def public_preview(job):
    return {k: job[k] for k in ("id", "filename", "version", "status", "warnings")} | {
        "sections": [{"location": s["location"], "text": s["text"][:500]} for s in job["sections"][:5]],
        "tables": [{"name": t["name"], "headers": t["headers"], "rows": t["rows"][:5], "row_count": len(t["rows"]), "suggested_mapping": suggestions(t)} for t in job["tables"]]}


def mapped_records(job, mappings):
    result = []
    if job["tables"] and len(mappings) != len(job["tables"]):
        raise ValueError("Confirm a mapping for every table")
    by_table = {m["table"]: m for m in mappings}
    all_ids = set()
    if len(by_table) != len(mappings):
        raise ValueError("Each table needs one mapping")
    for table in job["tables"]:
        mapping = by_table.get(table["name"])
        if not mapping or mapping["key_column"] not in table["headers"] or mapping["entity_type"] not in KINDS:
            raise ValueError("Invalid table mapping")
        if set(mapping["fields"]) - FIELDS or any(c not in table["headers"] for c in mapping["fields"].values()):
            raise ValueError("Unknown business field or column")
        seen = set()
        for number, row in enumerate(table["rows"], 2):
            record = dict(zip(table["headers"], row))
            key = record.get(mapping["key_column"], "").strip()
            if not key or key in seen:
                raise ValueError(f"Missing or duplicate key in {table['name']} row {number}")
            seen.add(key)
            fields = {f: record.get(c, "") for f, c in mapping["fields"].items()}
            for f in ("delivery_date", "required_date"):
                if fields.get(f):
                    fields[f] = date.fromisoformat(fields[f]).isoformat()
            for f in ("quantity", "stock"):
                if fields.get(f):
                    value = float(fields[f])
                    if not math.isfinite(value) or value < 0:
                        raise ValueError(f"{f} must be a finite non-negative number")
                    fields[f] = value
                elif f in fields:
                    del fields[f]
            kind = mapping["entity_type"]
            if kind + ":" + key in all_ids:
                raise ValueError("Duplicate business identifier across tables: " + kind + ":" + key)
            all_ids.add(kind + ":" + key)
            result.append({"id": kind + ":" + key, "entity_type": kind, "key": key,
                           "name": fields.get("name") or key, "fields": fields, "raw": record,
                           "location": table["name"] + "!A" + str(number), "active": True})
    return result


async def confirm(store: Store, job: dict, mappings: list[dict], router):
    if job["status"] == "completed":
        return job
    records = mapped_records(job, mappings)
    source_id = "source_" + hashlib.sha256(job["source_key"].encode()).hexdigest()[:20]
    old = await store.get("source", source_id)
    # Keep the detected event before any entity update so a retry cannot lose the change.
    changes = job.get("changes", [])
    if "changes" not in job:
        for record in records:
            prior = await store.get("entity", record["id"])
            if prior and record["entity_type"] == "PurchaseOrder" and prior["fields"].get("delivery_date") != record["fields"].get("delivery_date"):
                change_id = "change_" + hashlib.sha256((source_id + job["version"] + record["id"]).encode()).hexdigest()[:20]
                changes.append({"id": change_id, "entity_id": record["id"], "old_date": prior["fields"].get("delivery_date"), "new_date": record["fields"].get("delivery_date"), "title": "Supplier delivery date changed", "source_id": source_id, "version": job["version"], "status": "pending"})
    # Validate authority before changing any graph records.
    for record in records:
        prior = await store.get("entity", record["id"])
        if prior and prior.get("active") and prior["source_id"] != source_id and prior["fields"] != record["fields"]:
            raise ValueError(f"Conflicting source for {record['id']}. Resolve the authoritative record before importing.")
    if old and old["version"] == job["version"] and old.get("active") and old.get("import_complete"):
        job["status"] = "completed"
        job["source_id"] = source_id
        await store.put("import", job)
        return job
    job.update(changes=changes, status="running")
    await store.put("import", job)
    if old:
        for chunk in await store.list("chunk", 10000):
            if chunk["source_id"] == source_id:
                chunk["active"] = False
                await store.put("chunk", chunk)
                await store.graph.query(store.name, "MATCH (n:Chunk {id:$id}) REMOVE n.embedding, n.text", {"id": chunk["id"]})
        new_ids = {r["id"] for r in records}
        for entity in await store.list("entity", 10000):
            if entity["source_id"] == source_id and entity["id"] not in new_ids:
                entity["active"] = False
                await store.put("entity", entity)
    source = {"id": source_id, "name": job["filename"], "version": job["version"], "active": True,
              "record_count": len(records), "chunk_count": 0, "embedded_count": 0,
              "warnings": job["warnings"], "source_key": job["source_key"], "last_import_id": job["id"]}
    await store.put("source", source)
    version_id = source_id + ":" + job["version"]
    await store.put("source", {**source, "id": version_id, "is_version": True})
    if old:
        await store.link(version_id, "SUPERSEDES", source_id + ":" + old["version"])
    for record in records:
        record.update(source_id=source_id, version=job["version"])
        await store.put("entity", record)
        f = record["fields"]
        await store.graph.query(store.name, "MATCH (n:Entity {id:$id}) SET n.entity_type=$type, n.business_key=$key, n.name=$name, n.delivery_date=$delivery, n.required_date=$required, n.quantity=$quantity, n.stock=$stock, n.owner=$owner",
            {"id": record["id"], "type": record["entity_type"], "key": record["key"], "name": record["name"], "delivery": f.get("delivery_date", ""), "required": f.get("required_date", ""), "quantity": f.get("quantity", -1), "stock": f.get("stock", -1), "owner": f.get("owner", "")})
        await store.link(record["id"], "SUPPORTED_BY", version_id)
        text = json.dumps(record["raw"], ensure_ascii=False)
        job["sections"].append({"location": record["location"], "text": text, "entity_id": record["id"]})
    # Rebuild outbound dependencies from confirmed current records, including later-arriving tables.
    for record in await store.list("entity", 10000):
        if not record.get("active"):
            continue
        await store.graph.query(store.name, "MATCH (n:Entity {id:$id})-[r:DEPENDS_ON|OWNED_BY]->() DELETE r", {"id": record["id"]})
        f = record["fields"]
        for field, kind in (("supplier_id", "Supplier"), ("product_id", "Product"), ("purchase_order_id", "PurchaseOrder")):
            if f.get(field):
                await store.link(record["id"], "DEPENDS_ON", kind + ":" + f[field])
        if f.get("owner"):
            await store.link(record["id"], "OWNED_BY", "Team:" + f["owner"])
    degraded = False
    for section in job["sections"]:
        for i, text in enumerate(chunks(section["text"])):
            chunk = {"id": uid("chunk"), "source_id": source_id, "title": source["name"], "version": source["version"],
                     "location": section["location"] + (f" · part {i + 1}" if i else ""), "excerpt": text,
                     "entity_id": section.get("entity_id"), "active": True}
            await store.put("chunk", chunk)
            await store.graph.query(store.name, "MATCH (n:Chunk {id:$id}) SET n.text=$text", {"id": chunk["id"], "text": text})
            await store.link(source_id, "HAS_CHUNK", chunk["id"])
            await store.link(chunk["id"], "SUPPORTED_BY", version_id)
            if not degraded:
                try:
                    vector = await router.embed(text)
                    await store.graph.query(store.name, "MATCH (n:Chunk {id:$id}) SET n.embedding=vecf32($vector)", {"id": chunk["id"], "vector": vector})
                    source["embedded_count"] += 1
                except Exception:
                    degraded = True
            source["chunk_count"] += 1
    source["semantic_status"] = "ready" if source["embedded_count"] == source["chunk_count"] else "degraded"
    source["import_complete"] = True
    source["mappings"] = mappings
    if job.get("drive_file_id"):
        source.update(drive_file_id=job["drive_file_id"], drive_modified=job.get("drive_modified"))
    await store.put("source", source)
    for change in changes:
        await store.put("change", change)
    job.update(status="completed", source_id=source_id, changes=changes)
    await store.put("import", job)
    return job


async def revoke(store: Store, source_id: str):
    source = await store.get("source", source_id)
    if not source:
        raise ValueError("Source not found")
    source["active"] = False
    await store.put("source", source)
    for kind in ("chunk", "entity"):
        for record in await store.list(kind, 10000):
            if record.get("source_id") == source_id:
                record["active"] = False
                await store.put(kind, record)
                if kind == "chunk":
                    await store.graph.query(store.name, "MATCH (n:Chunk {id:$id}) REMOVE n.embedding, n.text", {"id": record["id"]})
